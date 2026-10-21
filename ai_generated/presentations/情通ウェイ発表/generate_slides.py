"""情通ウェイ発表資料（表紙 1 枚＋内容 1 枚）を CTC ブランドテンプレートから生成する。"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

W = "ai_generated/presentations/情通ウェイ発表"
NAVY = RGBColor(0x00, 0x3A, 0x85)
PALE = RGBColor(0xE8, 0xF0, 0xFA)
RED = RGBColor(0xD9, 0x2B, 0x0B)
TEXT = RGBColor(0x26, 0x26, 0x26)
LINE = RGBColor(0xC9, 0xD3, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

LEFT = [
    ("前進", "Physical AIへの挑戦",
     ["未経験の画像認識を自ら調査し、活用事例まで確認", "調査で終えず、実装・検証まで自分で実施"], NAVY),
    ("完遂", "WhiteLabel製品の品質向上",
     ["資料だけで判断せず、実機で動作・設定を確認", "パラメータシートを実機と細部まで突き合わせ"], NAVY),
    ("継承", "勉強会・ハンズオンによる知識共有",
     ["実機検証で得た知識を勉強会・ハンズオンで共有", "自分の中に留めず、周囲が使える状態にした"], RED),
]
RIGHT = [
    ("①", "判断や経験まで「形に残す」",
     ["成果物だけでなく「なぜその判断か」まで残す", "検証・試行錯誤で分かったことも知見にする"], RED),
    ("②", "自分だけの知識にしない",
     ["資料・パラメータシート・勉強会で周囲へ展開", "次の担当者がゼロから調べなくて済むようにする"], RED),
    ("③", "「起点」「完遂」の質を上げてから継承",
     ["起点：顧客やフロントが必要とすることまで考える", "完遂：品質を高めて最後までやり切り、次へ継承"], RED),
]

NOTES = """【1．これまで意識して行動できた取り組み】
1つ目は「前進」です。Physical AIへの取り組みで、これまで自分が経験したことのなかった画像認識などの技術について、技術そのものだけでなく、実際にどのように活用されているのかまで調査しました。また、調査だけで終わらせず、自分で手を動かして実装・検証するところまで進められたことは、「前進」を意識した取り組みだったと思っています。

2つ目は「完遂」です。WhiteLabelの業務では、製品資料を見るだけでなく、実際に実機を触って、設定や動作に問題がないかを細かく確認してきました。特にパラメータシートなど、今後も使われていく成果物については、実機と照らし合わせながら内容を確認し、一つひとつの仕事の品質を高めることを意識しました。

3つ目は「継承」です。WhiteLabelでは、自分自身が製品知識を身につけるだけでなく、勉強会やハンズオンを通じて、その知識や技術を周囲に伝える取り組みも行いました。自分が得た知識を自分の中だけに留めず、周囲の人も活用できる状態にするという部分は、「継承」につながる取り組みだったと思います。

【2．今後、最も推進したい情通ウェイ：継承】
今後、最も推進したいのは「継承」です。業務で作った成果物だけではなく、そこに至るまでに自分がどう考えたのか、なぜその判断をしたのかといった部分も、できるだけ形に残していきたいと思っています。そうすることで、今自分が経験していることが、将来誰かが同じ仕事をするときの価値になると考えています。

2つ目は、自分が得た知識を自分だけのものにしないことです。これまでの勉強会やハンズオンのように、自分が理解した内容を資料やナレッジとして周囲に展開して、次に同じ業務をする人がゼロから始めなくてもいい状態を作っていきたいです。

3つ目は、「継承」する内容そのものの質を上げることです。現状は、目の前の製品や技術を理解することに意識が向いてしまい、「お客様がなぜこれを必要としているのか」「何のためにこれを説明するのか」といった「起点」の部分まで考え切れていないことがあります。その結果、本当に学ぶべきポイントや、相手に伝えるべきポイントを十分に捉えられていないこともあると感じています。今後は、まず目的や背景をしっかり理解した上で、一つひとつの仕事の品質を高め、最後まで「完遂」する。そして、そこで得た経験や判断を次の人が使える形で「継承」するところまでを、一連の仕事として意識していきたいです。
"""


def rect(slide, shape, x, y, w, h, fill=None, line=None, line_w=0.75):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s


def put_text(shape, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(72000)
    tf.margin_top = tf.margin_bottom = Emu(36000)
    lines = text if isinstance(text, list) else [text]
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i:
            p.space_before = Pt(4)
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color


def textbox(slide, x, y, w, h, text, size, color, **kw):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    put_text(tb, text, size, color, **kw)
    return tb


def column(slide, x, w, head, cards, y0, card_h, gap):
    hd = rect(slide, MSO_SHAPE.RECTANGLE, x, y0, w, 420000, fill=NAVY)
    put_text(hd, head, 16, WHITE, bold=True)
    y = y0 + 420000 + gap
    for tag, title, bullets, accent in cards:
        rect(slide, MSO_SHAPE.RECTANGLE, x, y, w, card_h, fill=WHITE, line=LINE)
        rect(slide, MSO_SHAPE.RECTANGLE, x, y, 70000, card_h, fill=accent)
        chip = rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + 190000, y + 110000, 640000, 300000, fill=accent)
        put_text(chip, tag, 13, WHITE, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, x + 900000, y + 90000, w - 960000, 340000, title, 16, NAVY, bold=True)
        textbox(slide, x + 190000, y + 470000, w - 260000, card_h - 520000,
                ["・" + b for b in bullets], 14, TEXT, anchor=MSO_ANCHOR.TOP)
        y += card_h + gap
    return y


def drop(slide, idx):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == idx:
            ph._element.getparent().remove(ph._element)


def main():
    prs = Presentation(f"{W}/template.pptx")
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    layouts = {l.name: l for m in prs.slide_masters for l in m.slide_layouts}

    # 表紙
    s = prs.slides.add_slide(layouts["表紙（Blue）"])
    ph = {p.placeholder_format.idx: p for p in s.placeholders}
    ph[17].text_frame.text = "情通ウェイ共有"
    ph[0].text_frame.text = "情通ウェイの実践と、\n今後推進したい「継承」"
    ph[15].text_frame.text = "行動指針を意識した取り組みと、今後の展望"
    ph[18].text_frame.text = "2026年10月19日"
    drop(s, 13)
    drop(s, 14)

    # 内容（左右対比）
    s = prs.slides.add_slide(layouts["本文ページ（White）_2"])
    ph = {p.placeholder_format.idx: p for p in s.placeholders}
    ph[16].text_frame.text = "情通ウェイの実践と、今後推進する「継承」"
    drop(s, 15)
    drop(s, 17)

    cw, lx, rx = 5350000, 600000, 6242000
    y0, card_h, gap = 1130000, 1280000, 100000
    column(s, lx, cw, "【これまで】意識して行動できた取り組み", LEFT, y0, card_h, gap)
    column(s, rx, cw, "【今後】最も推進したいウェイ：継承", RIGHT, y0, card_h, gap)
    ay = y0 + 420000 + gap + (card_h * 3 + gap * 2) // 2 - 150000
    arrow = rect(s, MSO_SHAPE.RIGHT_ARROW, lx + cw + 18000, ay, 256000, 300000, fill=RED)
    arrow.name = "arrow"
    band_y = y0 + 420000 + gap + card_h * 3 + gap * 3
    band = rect(s, MSO_SHAPE.RECTANGLE, lx, band_y, rx + cw - lx, 400000, fill=PALE)
    put_text(band, "「起点」→「完遂」→「継承」を一連の仕事として意識し、経験を次の人へつなぐ",
             15, NAVY, bold=True, align=PP_ALIGN.CENTER)

    s.notes_slide.notes_text_frame.text = NOTES
    prs.save(f"{W}/output/情通ウェイ発表.pptx")


if __name__ == "__main__":
    main()
