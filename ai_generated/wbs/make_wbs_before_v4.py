import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

D=lambda m,d: dt.date(2026 if m>=10 else 2027,m,d)
START,END=D(10,21),D(2,28)
days=[START+dt.timedelta(i) for i in range((END-START).days+1)]
C0=13  # M列
FN="Meiryo UI"
thin=Side(style="thin",color="A6A6A6"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
HDR=PatternFill("solid",fgColor="D9E1F2")
f=lambda sz=10,b=False,c="000000": Font(name=FN,size=sz,bold=b,color=c)

wb=Workbook(); ws=wb.active; ws.title="WBS（日次）"
hol=wb.create_sheet("祝日"); hol["A1"]="日付"; hol["B1"]="名称"
for i,(d,n) in enumerate([(D(11,3),"文化の日"),(D(11,23),"勤労感謝の日"),(D(1,1),"元日"),(D(1,11),"成人の日"),(D(2,11),"建国記念の日"),(D(2,23),"天皇誕生日")],2):
    hol.cell(i,1,d).number_format="yyyy/m/d"; hol.cell(i,2,n)
hol["D1"]="日本の祝日のみ。会社の休業日(年末年始など)は必要に応じて追記してください。"
hol.column_dimensions["A"].width=14; hol.column_dimensions["B"].width=16

ws["B2"]="FY26US若手選抜出張_WBS"; ws["B2"].font=f(14,True); ws.merge_cells("B2:L2")
ws["M2"]="R：レビュー"; ws["M2"].font=f(10,True,"FF0000")
for c,t,rng in (("I3","計画","I3:J3"),("K3","実績","K3:L3")):
    ws[c]=t; ws[c].font=f(10,True); ws[c].alignment=Alignment(horizontal="center"); ws.merge_cells(rng)
heads=["カテゴリ","事務局","No","状況","タスク","担当","タスクの具体的作業(オプション)","開始日","終了日","開始日","終了日"]
for j,h in enumerate(heads):
    c=ws.cell(4,2+j,h); ws.merge_cells(start_row=4,start_column=2+j,end_row=5,end_column=2+j)
    c.font=f(10,True); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for r in (4,5): ws.cell(r,2+j).fill=HDR; ws.cell(r,2+j).border=BD
for k,d in enumerate(days):
    col=C0+k; c=ws.cell(5,col,d); c.number_format='d"日"'; c.font=f(8); c.alignment=Alignment(horizontal="center"); c.fill=HDR; c.border=BD
    ws.column_dimensions[L(col)].width=4.3
def runs(keyfn):
    out=[];s=0
    for k in range(1,len(days)+1):
        if k==len(days) or keyfn(days[k])!=keyfn(days[s]): out.append((s,k-1)); s=k
    return out
for s,e in runs(lambda d:(d.year,d.month)):
    c=ws.cell(4,C0+s,f"{days[s].month}月"); c.font=f(10,True); c.alignment=Alignment(horizontal="center")
    if e>s: ws.merge_cells(start_row=4,start_column=C0+s,end_row=4,end_column=C0+e)
    for k in range(s,e+1): ws.cell(4,C0+k).fill=HDR; ws.cell(4,C0+k).border=BD
mon=lambda d:d-dt.timedelta(d.weekday()); w0=mon(START)
ws["B6"]="週管理番号"; ws["B7"]="イベント"
for r in (6,7): ws.cell(r,2).font=f(9,True); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=12)
for s,e in runs(mon):
    c=ws.cell(6,C0+s,f"W{(mon(days[s])-w0).days//7+1}"); c.font=f(9); c.alignment=Alignment(horizontal="center")
    if e>s: ws.merge_cells(start_row=6,start_column=C0+s,end_row=6,end_column=C0+e)
col_of=lambda d:C0+(d-START).days
for d,t in {D(11,5):"KO",D(1,5):"出発",D(1,14):"帰国"}.items():
    c=ws.cell(7,col_of(d),t); c.font=f(7,True,"C00000"); c.alignment=Alignment(horizontal="center",shrink_to_fit=True)
for rr in (6,7):
    for col in range(C0,C0+len(days)): ws.cell(rr,col).border=BD

# ---- タスク定義 ---------------------------------------------------------
# (カテゴリ, 事務局区分, タスク, 担当, 具体的作業, 計画開始, 計画終了, 実績開始, 実績終了, 状況, 番号なし)
N=None; T=[]
def add(cat,sub,task,who="",note="",ps=N,pe=N,as_=N,ae=N,st="",child=False): T.append((cat,sub,task,who,note,ps,pe,as_,ae,st,child))
# 案内まで
add("案内まで","","予算獲得","","【実施済みなら日付を記入】")
add("案内まで","","企画書作成")
add("案内まで","","各本部へ企画書案内")
add("案内まで","","メンバー申込受付、確定","","確定リスト(追記版 1008)でUS希望7名を確定",as_=D(10,8),ae=D(10,8),st="完了")
# 訪問アレンジ
add("訪問アレンジ","","ラフな出張スケジュール立てる","","出張期間 1/5〜1/14(ラスベガス→サンフランシスコ→帰国)")
add("訪問アレンジ","","訪問先打診")
for n,t in (("Cisco","確定(サンフランシスコ)"),("Exaforce","確定(サンフランシスコ)"),("Forward Networks","確定(サンフランシスコ)"),("【訪問先D】","未確定"),("【訪問先E】","未確定"),("CTC America","確定(サンフランシスコ)")):
    add("訪問アレンジ","",f"●{n}","",f"【確定内容】\n・{t}\n・日時:\n【状況】",child=True)
add("訪問アレンジ","","CES 2027 参加登録","","2026/9/8時点で登録開始前。開始通知のメール設定済みか確認(CES参考情報)")
add("訪問アレンジ","","訪問先確定")
add("訪問アレンジ","","訪問先のコンテンツ検討/依頼")
add("訪問アレンジ","","訪問先の入館・見学に必要な書類/手続きの確認","","入館手続き、見学時の注意、参加者名簿の提出など。訪問先ごとに確認")
add("訪問アレンジ","","訪問先への事前連絡(アジェンダ・参加者)","","アジェンダ(時間割・移動)、参加者、言語を確認。【提案】訪問の1〜2週間前までに")
add("訪問アレンジ","","訪問先の頭出し","","訪問先の一覧と概要を、関係者へ先に提示する")
add("訪問アレンジ","","頭出しに沿ったパワポの作成","","訪問先の頭出しに沿って、訪問先紹介のPowerPointを作成")
add("訪問アレンジ","","スケジュール表の作成(詳細版)","","日別・時間帯別の細かいスケジュール表。移動・訪問・食事を含める")
# 出張前準備
S1="事前対応"
add("出張前準備",S1,"ホテルの予約(各自で予約)","","Mandalay Bay 1/5〜1/9。サンフランシスコ(1/9〜1/13)は【確定後に記入】")
add("出張前準備",S1,"ホテル内容ダブルチェック","","各自が予約した内容を集め、全員分の氏名と日付が合っているか確認。予約確認書の置き場所【調整】")
add("出張前準備",S1,"出張同行者リスト(e申請添付用)の作成・共有","","e申請に添付する同行者リスト。全員の氏名・本部・部署・社員番号・役職を記載。FY25のリストを参考")
add("出張前準備",S1,"入国手続き(ESTA・VISA等)の確認・案内","","【要確認】渡航先の要件を公式サイトで確認し、参加者へ案内。航空券・ホテルの予約時に必要になる場合あり")
add("出張前準備",S1,"TCP自主判定報告書の要否確認","","PC/スマホを持ち出し、データを渡航先へ供与しない場合は不要(2024/9/10変更)。供与する場合はe申請で申請。帰国後に申告番号を記載して再回付")
add("出張前準備",S1,"旅程作成・訪問先住所など込みで作成","","参加者ごとの旅程を一覧化(事務局管理)。一覧情報、旅のしおり、訪問先住所も記載。昨年度のものを参考")
add("出張前準備",S1,"出張先での食事・訪問先・移動手段の手配","","1か月前〜2週間前",D(12,4),D(12,18))
add("出張前準備",S1,"アンケート設問案作成","","Formsで昨年度と共通の最低限項目、改善リクエスト欄等を作成")
add("出張前準備",S1,"アンケート設問レビュー","","案をリーダー確認後、上長へレビュー依頼。●第1回: XX/XX XX:00〜XX:00 ●第2回: XX/XX XX:00〜XX:00")
add("出張前準備",S1,"事前費用の経費申請","","月をまたぐ場合は出発までに(またがない場合は帰国後でも可)")
add("出張前準備",S1,"FMCスマホの海外利用の事前設定","","1週間前",D(12,29),D(12,29))
S2="フライト・e申請・PC保険\n(事務局)"
add("出張前準備",S2,"フライトの候補を決める","","JTB-CWTに聞く。費用と空席を確認する")
add("出張前準備",S2,"JTB-CWTにフライト・海外旅行保険を手配する(事務局)","","承認済みの申請書+手配依頼書をメール送付。3か月前〜1か月前まで",D(10,21),D(12,4))
add("出張前準備",S2,"e申請(海外出張)(事務局)","","e申請(海外出張)・申請&承認。期限: 出発の1か月前(12/4)",D(10,21),D(12,4))
add("出張前準備",S2,"e申請内容をJTB-CWTに手配依頼書として送付","","事務局の申請書PDFを手配依頼書として送付。承認が下りたら速やかに連絡")
add("出張前準備",S2,"PC保険の依頼(事務局分)","","オリックス・レンテックから借用のPCは同社へ依頼(manual記載の宛先)。遅くとも出発1週間前(12/29)厳守。直前ガイダンスでリマインド",D(10,21),D(12,29))
S3="参加者対応"
add("出張前準備",S3,"出張者用のトカログループ作る","","参加者向けTocaroグループを作成し、初回案内を投稿")
add("出張前準備",S3,"↑トカロで、準備状況ノート作成・共有",child=True)
add("出張前準備",S3,"●e申請(海外出張)のアナウンス","","記載例(スクショ)、部課コードは各自で確認、同行者リスト添付、TOEIC点数(未受験は0)、役職区分、「課長確認不要」チェックの外し方、宿泊先名称・連絡先の書き方を案内",D(10,21),D(10,21))
add("出張前準備",S3,"●出張者へ各自JTBにフライトを手配をアナウンス","","フライト手配の方法や注意事項もアナウンスする【手配方法を調整】")
add("出張前準備",S3,"●PC保険申請をアナウンス","","SecurePC/シンクラ専用端末が対象。遅くとも出発1週間前(12/29)")
add("出張前準備",S3,"●ホテル予約のアナウンス(各自予約)","","Mandalay Bay 1/5〜1/9。早めの予約を依頼。予約期限【調整】",D(10,21),D(10,21))
add("出張前準備",S3,"●パスポート情報の送付依頼","","航空券予約に使用。期限・送付先【調整】")
add("出張前準備",S3,"(参加者)パスポート保有状況","","トカロ準備状況ノート確認・対応",child=True)
add("出張前準備",S3,"(参加者)パスポート新規申請者をフォロー",child=True)
add("出張前準備",S3,"(参加者)フライト予約・発券状況確認","","JTB-CWTへの予約依頼、発券状況を参加者ごとに確認",child=True)
add("出張前準備",S3,"(参加者)ホテル予約状況確認","","予約済みか、ホテル名・連絡先を確認",child=True)
add("出張前準備",S3,"(参加者)e申請の提出状況確認","","部課コードを各自で確認したうえで申請。期限: 出発の1か月前(12/4)",D(10,21),D(12,4),child=True)
add("出張前準備",S3,"(参加者)PC保険の申請状況確認","","遅くとも出発1週間前(12/29)",D(10,21),D(12,29),child=True)
S4="キックオフ"
add("出張前準備",S4,"キックオフ日程決定・アナウンス","","11/5(木) 17:00-18:00 神谷町(N11E)/オンライン可。会議室予約",D(10,21),D(10,21))
add("出張前準備",S4,"キックオフ前アナウンス","","【役員参加の有無を確認】。早めの集合を依頼。パスポート新規申請者に対する案内。【案】台湾版は開催の5日前",D(10,30),D(10,30))
add("出張前準備",S4,"キックオフ資料作成","","【案】昨年度資料をベースにFY26版を作成・更新。台湾版は開催の8日前〜2日前に作成(11/3は祝日)",D(10,28),D(11,2))
add("出張前準備",S4,"キックオフ資料へ「選出のお願い」資料を挿入","","「FY26情通事業G海外選抜研修選出のお願い r6.pptx」の2〜3枚目を、キックオフ資料に挿入")
add("出張前準備",S4,"キックオフ資料レビュー","","開催前に内容・役割分担を確認。【案】第1回 10/30、第2回 11/4")
add("出張前準備",S4,"キックオフ開催","","概要、スケジュール、手続き、懇親会を説明",D(11,5),D(11,5))
S6="CES予習"
add("出張前準備",S6,"CESの概要を見てもらう","","参加者にCESの概要(CES参考情報など)を共有。説明の場【キックオフ/直前ガイダンスで調整】")
add("出張前準備",S6,"セッションのピックアップ","","キーノートと一部のセッションは基本(参加登録)に含まれる。スタンダードでピックアップ。有料セッションを見たい場合は事務局へ相談")
add("出張前準備",S6,"ブース訪問先のピックアップ","","取り扱いのあるブースも可。新しいところも見られるよう、大手とスタートアップを混ぜて訪問")
add("出張前準備",S6,"注目テーマ3つを決める","","本部内で情報を収集(お客様の意向を含む)。幅広く見てもよい。3つのうち1つは、お客様や本部のテーマを反映したものにする")
S5="直前ガイダンス"
add("出張前準備",S5,"直前ガイダンスとレビュー(2回)の日程調整","","それぞれ1.5時間程度。ガイダンス1回+レビュー2回")
add("出張前準備",S5,"直前ガイダンス日程・会場確定","","台湾版は出発の14日前に開催(US は12/22頃に相当。年末のため要調整)")
add("出張前準備",S5,"直前ガイダンス資料作成","","昨年資料から更新:出張前ToDo、細かい手続き(保険、eSIMなど)、Uberなど現地移動、経費精算、情通の取組み、アンケート協力依頼等。入国手続き【要確認】")
add("出張前準備",S5,"直前ガイダンス説明担当・内容調整")
add("出張前準備",S5,"直前ガイダンス資料レビュー(第1回)","","●XX/XX XX:00〜XX:30(1.5時間程度)",child=True)
add("出張前準備",S5,"直前ガイダンス資料レビュー(第2回)","","●XX/XX XX:00〜XX:30(1.5時間程度)",child=True)
add("出張前準備",S5,"直前ガイダンス開催","","参加者へ出張前の最終案内を実施。PC保険の期限(1週間前)をリマインド")
add("出張前準備",S5,"細かい手続きの案内(保険・eSIMなど)","","海外旅行保険、PC保険、eSIM(海外通信)の手続きを案内")
add("出張前準備",S5,"Uber(配車アプリ)の案内","","出発前にUberを入れてもらう。昨年度資料を参考に、利用時の注意事項を案内")
add("出張前準備",S5,"アンケート回答に関する事前案内","","帰国直後に実施し、結果を上司・本部長等へ共有する可能性を案内")
add("出張前準備",S5,"持ち物・服装の案内","","【渡航先(米国・1月)に合わせた持ち物・服装】")
add("出張前準備",S5,"緊急連絡先の整理・共有","","参加者が確認できる場所へ掲載")
# 出張中
add("出張中","","出張へ行く(US出張)","参加者・事務局","出張期間 1/5〜1/14",D(1,5),D(1,14))
add("出張中","","出国・ラスベガス(CES)","","NH008 成田→サンフランシスコ→UA661 ラスベガス。宿泊: Mandalay Bay。CES 3日間の日付【確認】",D(1,5),D(1,9))
add("出張中","","サンフランシスコ(ベンダー訪問)","","Cisco / Exaforce / Forward Networks / CTC America ほか。宿泊先【調整】",D(1,9),D(1,13))
add("出張中","","帰国","","NH107 サンフランシスコ→羽田(1/14 羽田着)",D(1,13),D(1,14))
add("出張中","","集合・移動・訪問の進行確認","","当日の集合、移動、訪問スケジュールを確認",D(1,5),D(1,14))
add("出張中","","参加者の連絡・トラブル対応","","緊急時は整理した連絡先を使用",D(1,5),D(1,14))
add("出張中","","報告資料用の情報収集","","訪問先、学び、所感、アクションプランを記録",D(1,5),D(1,14))
# 帰国後
add("帰国後","","アンケート配信","","帰国後すぐに参加者へ配信",D(1,14),D(1,15))
add("帰国後","","アンケート回答回収・未回答者フォロー","","回収期限を設定し、未回答者をフォロー")
add("帰国後","","アンケート集計・サマリー作成","","満足度、有意義だった内容、改善要望等を集計")
add("帰国後","","TCP自主判定報告書の再回付(該当時のみ)","","TCP申請した場合のみ。帰国後1週間以内に申告番号を記載して再回付",D(1,14),D(1,21))
add("帰国後","","訪問先へのお礼連絡","","【提案】訪問先ごとに、帰国後すぐに連絡。質疑で出た宿題があれば回答")
add("帰国後","","海外出張申請書の再回付","","帰国後1週間以内(旅程変更がなくても必要)",D(1,14),D(1,21))
add("帰国後","","旅費精算(日当計算・経費精算)","","規程: 帰国後1週間以内。マニュアルは月末と記載のため要確認",D(1,14),D(1,21))
add("帰国後","","出張報告書テンプレート準備","","FY25台湾の報告書を参考に作成済み。訪問先・日程の反映が残り",D(10,19),D(10,31))
add("帰国後","","報告書を作成する","参加者・事務局","参加者版、事務局版、本部長向け報告資料を作成")
add("帰国後","","｜出張報告書(Word)作成・提出","","マニュアルに「2週間以内」「1か月以内」の記載が混在。要確認",D(1,15),D(1,28),child=True)
add("帰国後","","｜出張報告書(PowerPoint)作成","","報告会当日までに作成。メンバーで分担【調整】",child=True)
add("帰国後","","｜(参加者)報告書確認",child=True)
add("帰国後","","｜(事務局)報告書作成",child=True)
add("帰国後","","｜レビュー(上長)",child=True)
add("帰国後","","｜本部長向け報告資料の作成","","参加メンバーに協力を依頼し、リーダー監修で完成させる",child=True)
add("帰国後","","｜本部長向け報告資料の上長レビュー","","送付前にレビュー依頼",child=True)
add("帰国後","","出張報告会","","開催日【調整】")
add("帰国後","","C-WE掲載版の作成","","昨年度掲載版を参考に作成")
add("帰国後","","経費精算の実施・フォロー","","台湾版の目安: 帰国翌週(US は1/18週)、遅くとも翌々週(1/25週)までに完了")
add("帰国後","","成果物・運営資料の格納","","FY26 Boxフォルダへ資料、アンケート、報告書等を整理して保存")
add("帰国後","","振り返り・次年度引継ぎ事項整理","","改善点、未完了事項、次年度向けの引継ぎを記録")

for _ in range(25): add('', '', '', '', '', child=True)
r0=8; r=r0; no=0; groups=[]; subs=[]
for cat,sub,task,who,note,ps,pe,as_,ae,st,child in T:
    if not child: no+=1
    ws.cell(r,4,None if child else no); ws.cell(r,5,st); ws.cell(r,6,task); ws.cell(r,7,who); ws.cell(r,8,note)
    for j,v in enumerate((ps,pe,as_,ae)): ws.cell(r,9+j,v).number_format="m/d"
    if cat:
        if not groups or groups[-1][0]!=cat: groups.append([cat,r,r])
        else: groups[-1][2]=r
    if sub:
        if not subs or subs[-1][0]!=sub or subs[-1][2]!=r-1: subs.append([sub,r,r])
        else: subs[-1][2]=r
    for col in range(2,C0+len(days)):
        c=ws.cell(r,col); c.border=BD
        if col<C0: c.font=f(10); c.alignment=Alignment(vertical="center",wrap_text=(col in (6,8)),horizontal="center" if col in (4,5,9,10,11,12) else None)
    ws.row_dimensions[r].height=max(15,14*max(note.count("\n")+1,1+len(note)//38,1+len(task)//22))
    r+=1
last=r-1
for lst,col in ((groups,2),(subs,3)):
    for name,a,b in lst:
        c=ws.cell(a,col,name); c.font=f(10,True); c.alignment=Alignment(vertical="center",wrap_text=True)
        if b>a: ws.merge_cells(start_row=a,start_column=col,end_row=b,end_column=col)
dv=DataValidation(type="list",formula1='"未着手,進行中,完了"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"E{r0}:E{last}")
for k,w in {"A":3.8,"B":11,"C":16,"D":4.5,"E":8.5,"F":40,"G":12,"H":56,"I":7.5,"J":7.5,"K":7.5,"L":7.5}.items(): ws.column_dimensions[k].width=w
rng=f"M{r0}:{L(C0+len(days)-1)}{last}"
dx=lambda color: DifferentialStyle(fill=PatternFill(bgColor=color,fill_type="solid"))
rule=lambda fm,color: Rule(type="expression",formula=[fm],dxf=dx(color))
ws.conditional_formatting.add(rng,rule('AND($K8<>"",$L8<>"",M$5>=$K8,M$5<=$L8)',"9DC3E6"))
ws.conditional_formatting.add(rng,rule('AND($I8<>"",$J8<>"",M$5>=$I8,M$5<=$J8)',"4472C4"))
ws.conditional_formatting.add(rng,rule('OR(WEEKDAY(M$5,2)>5,COUNTIF(祝日!$A:$A,M$5)>0)',"FCE4D6"))
ws.conditional_formatting.add(f"M5:{L(C0+len(days)-1)}5",rule('OR(WEEKDAY(M$5,2)>5,COUNTIF(祝日!$A:$A,M$5)>0)',"F8CBAD"))
help_=wb.create_sheet("使い方",0)
help_.column_dimensions["A"].width=3; help_.column_dimensions["B"].width=22; help_.column_dimensions["C"].width=100
rows=[("このWBSについて","FY26 US若手選抜出張(2027/1/5〜1/14)の準備・実施・報告のタスク管理表です。台湾版WBSを参考に作成しています。"),
("書き方",""),
("  タスクの追加","一番下の空行に入力します。カテゴリ・事務局の列は、上の行からコピーして貼り付けてください。"),
("  計画・実績の日付","開始日・終了日を入力すると、右側のガントにバーが自動で出ます(計画=濃い青、実績=薄い青)。1日だけの作業は、開始日と終了日に同じ日付を入れます。"),
("  状況","プルダウンから、未着手・進行中・完了を選びます。"),
("  担当","担当者名を入力します。"),
("  タスクの具体的作業","補足、期限の根拠、メモを書きます。【調整】は、未確定の項目です。【案】は、台湾版などを参考にした仮の日程です。"),
("色",""),
("  オレンジ","土日と祝日です。会社の休業日は、「祝日」シートに追記すると反映されます。"),
("  KO / 出発 / 帰国","イベント行(7行目)に表示しています。"),
("レビュー","R はレビューの印です。レビューがあるタスクの日付の列に入力してください。"),
("注意","日付は、2026/10/21〜2027/2/28の範囲で入力してください。範囲外はバーが出ません。")]
for i,(a,b) in enumerate(rows,2):
    help_.cell(i,2,a).font=f(10,not a.startswith("  ")); help_.cell(i,3,b).font=f(10); help_.cell(i,3).alignment=Alignment(wrap_text=True,vertical="top")
wb.active=1
dvd=DataValidation(type="date",operator="between",formula1="DATE(2026,10,1)",formula2="DATE(2027,12,31)",allow_blank=True,showErrorMessage=True,errorTitle="日付",error="日付(例: 11/5)を入力してください。")
ws.add_data_validation(dvd); dvd.add(f"I{r0}:L{last}")
ws.freeze_panes="A8"; ws.sheet_view.zoomScale=85
ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=8
out="ai_generated/wbs/FY26US若手選抜出張_WBS_枠完成版.xlsx"; wb.save(out); print(out,"最終行",last,"番号付き",no)
