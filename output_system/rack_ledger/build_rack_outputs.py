#!/usr/bin/env python3
"""機器台帳（rack_ledger.xlsx）から、ラック搭載図と電力集計を作り直す。

台帳が正本。図（drawio）と集計（xlsx）は毎回捨てて作り直す前提の生成物。
"""
import argparse
import html
import re
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

U_H = 20          # 1U の高さ(px)
INNER_W = 330     # ラック内寸の幅(px)
ULABEL_W = 26
PAD = 40
GAP = 90          # 前面図と背面図の間隔

USAGE_COLOR = {
    "Server": "#DAE8FC", "Storage": "#FFE6CC", "SW": "#D5E8D4",
    "PDU": "#F8CECC", "FW": "#E1D5E7", "LB": "#E1D5E7",
}
POWER_OFF = re.compile(r"電源オフ|電源未投入|未投入|電源断|停止中|未使用|撤去")
NON_POWERED = re.compile(r"ｹｰﾌﾞﾙ|ケーブル|PatchPnael|パッチ|使用不可|ﾌﾞﾗﾝｸ|ブランク|棚板|棚|"
                         r"ｱﾝｸﾞﾙ|アングル|引き出し|引出|ﾏﾈｰｼﾞﾒﾝﾄ")


# mxCell の value は XML 属性なので、HTML タグもエスケープした形で入れる
BR = "&lt;br&gt;"


def esc(s):
    return html.escape(str(s if s is not None else ""), quote=True)


def small(s):
    return f"&lt;font style=&quot;font-size:8px&quot;&gt;{esc(s)}&lt;/font&gt;"


def num(v):
    return v if isinstance(v, (int, float)) else None


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["機器台帳"]
    heads = [c.value for c in ws[1]]
    devices = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(heads, r))
        if not rec.get("ラック") or not rec.get("機器名"):
            continue
        rec["_circuits"] = []
        for i in (1, 2, 3):
            lb = rec.get(f"回路{i}")
            cur = rec.get(f"電流{i}A")
            if lb and num(cur) is not None:
                rec["_circuits"].append({"label": str(lb).strip(), "current": float(cur),
                                         "ports": num(rec.get(f"口数{i}"))})
        devices.append(rec)
    ws2 = wb["ラック定義"]
    h2 = [c.value for c in ws2[1]]
    racks = {}
    for r in ws2.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(h2, r))
        if not rec.get("ラック"):
            continue
        rk = racks.setdefault(rec["ラック"], {
            "name": rec["ラック"], "sheet": rec.get("シート"), "spec": rec.get("寸法"),
            "total_u": int(rec.get("総U") or 47), "circuits": [],
            "demand": float(rec.get("需要率") or 0.6), "warn": float(rec.get("警告閾値") or 0.8)})
        if rec.get("回路"):
            rk["circuits"].append({"label": str(rec["回路"]).strip(),
                                   "volt": num(rec.get("電圧V")),
                                   "breaker": num(rec.get("ブレーカA"))})
    return racks, devices


def lanes_for(items):
    """U範囲が重なる機器を横に並べるため、レーン番号と同じ塊のレーン数を返す"""
    items = sorted(items, key=lambda d: (-(d["開始U"] + d["U数"] - 1), d["機器名"] or ""))
    comp, cur, cur_lo = [], [], None
    for d in items:
        top = d["開始U"] + d["U数"] - 1
        if cur and top < cur_lo:
            comp.append(cur)
            cur, cur_lo = [], None
        cur.append(d)
        cur_lo = d["開始U"] if cur_lo is None else min(cur_lo, d["開始U"])
    if cur:
        comp.append(cur)
    out = {}
    for group in comp:
        lanes = []  # lanes[i] = 最小U（そのレーンで使用済みの下端）
        for d in group:
            top = d["開始U"] + d["U数"] - 1
            placed = False
            for li, lo in enumerate(lanes):
                if top < lo:
                    lanes[li] = d["開始U"]
                    out[id(d)] = li
                    placed = True
                    break
            if not placed:
                lanes.append(d["開始U"])
                out[id(d)] = len(lanes) - 1
        for d in group:
            out[(id(d), "n")] = len(lanes)
    return out


def power(rack, devs):
    """回路ごとの定格・両系・片系と判定を返す"""
    res = []
    for ci in rack["circuits"]:
        rated = 0.0
        both = 0.0
        for d in devs:
            mine = [c for c in d["_circuits"] if c["label"] == ci["label"]]
            if not mine:
                continue
            n = len(d["_circuits"])
            for c in mine:
                rated += c["current"]
                both += c["current"] / n
        single_eff = rated * rack["demand"]
        both_eff = both * rack["demand"]
        limit = (ci["breaker"] or 0) * rack["warn"]
        ratio = single_eff / ci["breaker"] if ci["breaker"] else 0
        res.append({
            "label": ci["label"], "volt": ci["volt"], "breaker": ci["breaker"],
            "rated": rated, "both_eff": both_eff, "single_eff": single_eff,
            "limit": limit, "ratio": ratio,
            "judge": "要対応（片系で閾値超過）" if ci["breaker"] and single_eff > limit else "余裕あり",
        })
    watt = sum((num(d.get("電圧V")) or 0) * c["current"] for d in devs for c in d["_circuits"])
    weight = sum(num(d.get("重量kg")) or 0 for d in devs)
    return res, watt, weight


def elevation(rack, devs, face, x0, y0):
    """前面図または背面図の mxCell 群を返す"""
    cells = []
    tu = rack["total_u"]
    h = tu * U_H
    cells.append(f'<mxCell id="{face}_title_{rack["name"]}" value="{esc(rack["name"])} {esc(face)}" '
                 f'style="text;html=1;fontSize=14;fontStyle=1;align=center;" vertex="1" parent="1">'
                 f'<mxGeometry x="{x0}" y="{y0 - 30}" width="{ULABEL_W + INNER_W}" height="24" as="geometry"/></mxCell>')
    cells.append(f'<mxCell id="{face}_frame_{rack["name"]}" value="" '
                 f'style="rounded=0;html=1;fillColor=none;strokeColor=#333333;strokeWidth=2;" vertex="1" parent="1">'
                 f'<mxGeometry x="{x0 + ULABEL_W}" y="{y0}" width="{INNER_W}" height="{h}" as="geometry"/></mxCell>')
    for u in range(1, tu + 1):
        y = y0 + (tu - u) * U_H
        cells.append(f'<mxCell id="{face}_u{u}_{rack["name"]}" value="{u}" '
                     f'style="text;html=1;fontSize=8;align=right;verticalAlign=middle;strokeColor=none;fillColor=none;" '
                     f'vertex="1" parent="1"><mxGeometry x="{x0}" y="{y}" width="{ULABEL_W - 3}" '
                     f'height="{U_H}" as="geometry"/></mxCell>')
        cells.append(f'<mxCell id="{face}_g{u}_{rack["name"]}" value="" '
                     f'style="line;html=1;strokeColor=#DDDDDD;" vertex="1" parent="1">'
                     f'<mxGeometry x="{x0 + ULABEL_W}" y="{y}" width="{INNER_W}" height="1" as="geometry"/></mxCell>')

    mine = [d for d in devs if d.get("面") in (face, "両面")]
    lane = lanes_for(mine)
    for i, d in enumerate(mine):
        top = d["開始U"] + d["U数"] - 1
        y = y0 + (tu - top) * U_H
        hh = d["U数"] * U_H
        n = lane[(id(d), "n")]
        li = lane[id(d)]
        w = INNER_W / n
        x = x0 + ULABEL_W + li * w
        name = str(d["機器名"]).strip()
        powered = not NON_POWERED.search(name)
        fill = USAGE_COLOR.get(str(d.get("用途") or "").strip(), "#F5F5F5") if powered else "#EEEEEE"
        off = bool(POWER_OFF.search(str(d.get("備考") or "")))
        missing = powered and not d["_circuits"] and not off
        # 1U は高さ 20px しかないため1行に収める。2U以上は改行して読みやすくする
        parts = [esc(name)]
        # 横に並べて幅が狭いときは資産番号を省く（枠からはみ出すため）
        if d.get("資産番号") and w >= 150:
            parts.append(small(d["資産番号"]))
        if missing:
            parts.append(small("電力未記入"))
        elif off:
            parts.append(small("電源オフ"))
        sep = " " if d["U数"] == 1 else BR
        label = sep.join(parts)
        if missing:
            stroke = "strokeColor=#CC0000;dashed=1;strokeWidth=2;"
        elif off:
            stroke = "strokeColor=#999999;dashed=1;"
        else:
            stroke = "strokeColor=#666666;"
        fs = 8 if d["U数"] == 1 else 9
        cells.append(
            f'<mxCell id="{face}_d{i}_{rack["name"]}" value="{label}" '
            f'style="rounded=0;html=1;whiteSpace=wrap;overflow=hidden;verticalAlign=middle;'
            f'fillColor={fill};{stroke}fontSize={fs};" '
            f'vertex="1" parent="1"><mxGeometry x="{x:.0f}" y="{y}" width="{w:.0f}" height="{hh}" as="geometry"/></mxCell>')
    return cells, h


def drawio_for(rack, devs, pw, watt, weight):
    tu = rack["total_u"]
    y0 = PAD + 40
    front, h = elevation(rack, devs, "前面", PAD, y0)
    back, _ = elevation(rack, devs, "背面", PAD + ULABEL_W + INNER_W + GAP, y0)
    cells = front + back
    lines = [f"寸法: {rack['spec'] or ''}　総U: {tu}U　機器総重量: {weight:g} kg",
             f"定格消費電力（回路別定格の合計）: {watt:g} W　"
             f"予測実効値（需要率 {rack['demand']:g}）: {watt * rack['demand']:g} W"]
    for p in pw:
        lines.append(
            f"{p['label']}: 定格 {p['rated']:g}A / 両系実効 {p['both_eff']:.2f}A / "
            f"片系実効 {p['single_eff']:.2f}A → ブレーカ {p['breaker']:g}A の "
            f"{p['ratio']:.0%}（閾値 {rack['warn']:.0%}）… {p['judge']}")
    lines.append("凡例: 青=Server / 橙=Storage / 緑=SW / 赤=PDU / 灰=非電源品　"
                 "赤い破線＝電力が未記入で集計に入っていない機器　"
                 "灰の破線＝備考で電源オフとされている機器")
    txt = BR.join(esc(x) for x in lines)
    sw = (ULABEL_W + INNER_W) * 2 + GAP
    cells.append(f'<mxCell id="sum_{rack["name"]}" value="{txt}" '
                 f'style="rounded=0;html=1;whiteSpace=wrap;align=left;verticalAlign=top;fillColor=#FFF7E6;'
                 f'strokeColor=#D6B656;fontSize=10;spacing=6;" vertex="1" parent="1">'
                 f'<mxGeometry x="{PAD}" y="{y0 + h + 24}" width="{sw}" height="{28 + 16 * len(lines)}" as="geometry"/></mxCell>')
    body = "\n".join(cells)
    return (f'<mxfile host="gaido"><diagram name="{esc(rack["name"])}">'
            f'<mxGraphModel dx="800" dy="600" grid="0" gridSize="10" guides="1" tooltips="1" '
            f'connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="1169" pageHeight="826" '
            f'math="0" shadow="0"><root><mxCell id="0"/><mxCell id="1" parent="0"/>\n{body}\n'
            f'</root></mxGraphModel></diagram></mxfile>')


def check(rack, devs):
    """台帳そのものの点検"""
    out = []
    for face in ("前面", "背面"):
        occ = defaultdict(list)
        for d in devs:
            if d.get("面") not in (face, "両面"):
                continue
            if NON_POWERED.search(str(d["機器名"])):
                continue
            for u in range(d["開始U"], d["開始U"] + d["U数"]):
                occ[u].append(d["機器名"])
        for u, names in sorted(occ.items()):
            if len(names) > 1:
                out.append(f"(参考) {rack['name']} {face} {u}U に {len(names)}台: {' / '.join(map(str, names))}")
    for d in devs:
        top = d["開始U"] + d["U数"] - 1
        if d["開始U"] < 1 or top > rack["total_u"]:
            out.append(f"要対応 {rack['name']} {d['機器名']}: {d['開始U']}U から {d['U数']}U は "
                       f"総 {rack['total_u']}U の外にはみ出す")
        if NON_POWERED.search(str(d["機器名"])):
            continue
        if not d["_circuits"]:
            if POWER_OFF.search(str(d.get("備考") or "")):
                out.append(f"(参考) {rack['name']} {d['開始U']}U {d['機器名']}: "
                           f"電力の記入なし（備考『{str(d.get('備考')).strip()}』のため意図的とみなす）")
            else:
                out.append(f"要対応 {rack['name']} {d['開始U']}U {d['機器名']}: 電力が未記入（集計に入らない）")
        if num(d.get("重量kg")) is None:
            out.append(f"要対応 {rack['name']} {d['開始U']}U {d['機器名']}: 重量が未記入")
        for k in ("電圧V", "系統数", "重量kg", "電流1A", "電流2A", "電流3A", "口数1", "口数2", "口数3"):
            v = d.get(k)
            if isinstance(v, str) and v.strip():
                out.append(f"要対応 {rack['name']} {d['開始U']}U {d['機器名']}: {k} が文字列『{v}』")
        known = {c["label"] for c in rack["circuits"]}
        for c in d["_circuits"]:
            if c["label"] not in known:
                out.append(f"要対応 {rack['name']} {d['開始U']}U {d['機器名']}: "
                           f"回路『{c['label']}』が『ラック定義』に無い")
    return out


def write_summary(path, rows, checks):
    wb = Workbook()
    ws = wb.active
    ws.title = "電力集計"
    heads = ["ラック", "回路", "電圧V", "ブレーカA", "定格電流A", "両系実効A", "片系実効A",
             "片系/ブレーカ", "警告閾値", "判定"]
    ws.append(heads)
    for i, h in enumerate(heads, 1):
        ws.cell(1, i).font = Font(bold=True)
        ws.cell(1, i).fill = PatternFill("solid", fgColor="DDEBF7")
        ws.cell(1, i).alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(h) + 4)
    ng = PatternFill("solid", fgColor="FFC7CE")
    for r in rows:
        ws.append(r)
        if str(r[-1]).startswith("要対応"):
            for c in range(1, len(heads) + 1):
                ws.cell(ws.max_row, c).fill = ng
    ws.cell(1, 8).number_format = "0%"
    for rr in range(2, ws.max_row + 1):
        ws.cell(rr, 8).number_format = "0%"
        ws.cell(rr, 9).number_format = "0%"
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("点検結果")
    ws2.append(["区分", "内容"])
    for i in (1, 2):
        ws2.cell(1, i).font = Font(bold=True)
        ws2.cell(1, i).fill = PatternFill("solid", fgColor="DDEBF7")
    ws2.column_dimensions["A"].width = 10
    ws2.column_dimensions["B"].width = 110
    for c in checks:
        kind, _, body = c.partition(" ")
        ws2.append([kind, body])
        if kind == "要対応":
            ws2.cell(ws2.max_row, 1).fill = ng
    ws2.freeze_panes = "A2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ledger", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--png", action="store_true", help="drawio を PNG へ書き出す")
    a = ap.parse_args()

    racks, devices = load(a.ledger)
    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    by_rack = defaultdict(list)
    for d in devices:
        by_rack[d["ラック"]].append(d)

    rows, checks, made = [], [], []
    for name, rack in racks.items():
        devs = by_rack.get(name, [])
        pw, watt, weight = power(rack, devs)
        for p in pw:
            rows.append([name, p["label"], p["volt"], p["breaker"], round(p["rated"], 2),
                         round(p["both_eff"], 2), round(p["single_eff"], 2),
                         round(p["ratio"], 4), rack["warn"], p["judge"]])
        checks += check(rack, devs)
        f = out / f"{name}_搭載図.drawio"
        f.write_text(drawio_for(rack, devs, pw, watt, weight), encoding="utf8")
        made.append(f)
        print(f"{name}: 機器{len(devs)}件 定格{watt:g}W 重量{weight:g}kg → {f.name}", file=sys.stderr)

    write_summary(out / "power_summary.xlsx", rows, checks)
    print(f"電力集計: {out / 'power_summary.xlsx'}", file=sys.stderr)
    ng = [c for c in checks if c.startswith("要対応")]
    print(f"点検: 要対応 {len(ng)}件 / 参考 {len(checks) - len(ng)}件", file=sys.stderr)

    if a.png:
        tool = Path(__file__).resolve().parents[2] / "tools" / "drawio_export.py"
        for f in made:
            png = f.with_suffix(".png")
            r = subprocess.run([sys.executable, str(tool), str(f), str(png), "--scale", "2"],
                               capture_output=True, text=True)
            print(f"  PNG {png.name}: exit={r.returncode} {r.stderr.strip()[-200:]}", file=sys.stderr)


if __name__ == "__main__":
    main()
