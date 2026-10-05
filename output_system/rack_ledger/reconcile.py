#!/usr/bin/env python3
"""原本 Excel の集計値と、台帳から再計算した値を突き合わせる。

差が出た原因になっているセルを列挙する。移行時と、原本を並行運用している間の検算に使う。
"""
import argparse
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from extract_ledger import resolve_columns, find_bands
from build_rack_outputs import load, power


def orig_totals(xlsx):
    """原本の『システムラック総電気容量（定格）』『機器総重量』を読む"""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out = {}
    for ws in wb.worksheets:
        name = str(ws.cell(1, 1).value).strip()
        rec = {"watt": None, "weight": None, "watt_cell": "", "weight_cell": ""}
        bands = find_bands(ws)
        lo = bands[0][2] + 1
        for r in range(lo, lo + 12):
            lab = str(ws.cell(r, 10).value or "")
            if not lab:
                continue
            for c in range(19, 32):
                v = ws.cell(r, c).value
                if not isinstance(v, (int, float)):
                    continue
                if "総電気容量（定格）" in lab and rec["watt"] is None and v > 100:
                    rec["watt"], rec["watt_cell"] = v, ws.cell(r, c).coordinate
                elif "機器総重量" in lab and rec["weight"] is None:
                    rec["weight"], rec["weight_cell"] = v, ws.cell(r, c).coordinate
        # ラベルが1行ずれているシートは、重量が『電流値計（定格）』行にある
        if rec["weight"] is None:
            for r in range(lo, lo + 12):
                v = ws.cell(r, 20).value
                if isinstance(v, (int, float)) and 10 < v < 1000:
                    rec["weight"], rec["weight_cell"] = v, ws.cell(r, 20).coordinate
                    break
        out[name] = rec
    return out


def causes(xlsx):
    """手入力の消費電力(W) と 電圧×電流 が食い違う行を挙げる"""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out = {}
    for ws in wb.worksheets:
        name = str(ws.cell(1, 1).value).strip()
        cir, cols = resolve_columns(ws)
        rows = []
        for face, top, bot in find_bands(ws):
            for r in range(top, bot + 1):
                ac = ws.cell(r, cols["watt"]).value if cols["watt"] else None
                v = ws.cell(r, cols["volt"]).value
                cur = [ws.cell(r, c["col_current"]).value for c in cir]
                nums = [x for x in cur if isinstance(x, (int, float))]
                vi = v * sum(nums) if isinstance(v, (int, float)) and nums else None
                ac_n = ac if isinstance(ac, (int, float)) else None
                if ac_n is None and vi is None:
                    continue
                if ac_n is None:
                    rows.append(f"{ws.title}!行{r}（{int(ws.cell(r, 3).value)}U）: "
                                f"消費電力が未記入。電圧×電流 = {vi:g}W "
                                f"→ 原本の合計には入っていない")
                elif vi is None:
                    rows.append(f"{ws.title}!行{r}（{int(ws.cell(r, 3).value)}U）: "
                                f"手入力 {ac_n:g}W だが電圧か電流が数値でなく検算できない")
                elif abs(ac_n - vi) > 0.5:
                    rows.append(f"{ws.title}!行{r}（{int(ws.cell(r, 3).value)}U）: "
                                f"手入力 {ac_n:g}W ≠ 電圧{v:g}V×電流計{sum(nums):g}A = {vi:g}W")
        out[name] = rows
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--original", required=True)
    ap.add_argument("--ledger", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    ot = orig_totals(a.original)
    cz = causes(a.original)
    racks, devices = load(a.ledger)
    by = {}
    for d in devices:
        by.setdefault(d["ラック"], []).append(d)

    lines = ["# 原本と台帳の突合", "",
             "原本の集計値と、台帳から計算し直した値を比べます。差が残っている行は、"
             "原本の記入に判断が必要なものです。", "",
             "| ラック | 原本の定格W | 台帳の定格W | 差 | 原本の重量kg | 台帳の重量kg | 差 |",
             "|---|---:|---:|---:|---:|---:|---:|"]
    detail = []
    for name, rack in racks.items():
        devs = by.get(name, [])
        pw, watt, weight = power(rack, devs)
        o = ot.get(name, {})
        ow, og = o.get("watt"), o.get("weight")
        dw = f"{watt - ow:+g}" if isinstance(ow, (int, float)) else "—"
        dg = f"{weight - og:+g}" if isinstance(og, (int, float)) else "—"
        lines.append(f"| {name} | {ow if ow is not None else '—'} | {watt:g} | {dw} | "
                     f"{og if og is not None else '—'} | {weight:g} | {dg} |")
        if cz.get(name):
            detail.append(f"## {name} で差の原因になっている行")
            detail.append("")
            detail += [f"- {x}" for x in cz[name]]
            detail.append("")
    lines += ["", "差が 0 のラックは、台帳が原本と同じ値を再現できています。", ""] + detail
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text("\n".join(lines), encoding="utf8")
    print(f"突合結果: {a.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
