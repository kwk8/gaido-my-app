#!/usr/bin/env python3
"""LinkSphere パラメーターシートを、Huawei（SphereOS）の用語と既定値で作り直す。

元のブックから「改訂履歴」シートだけを引き継ぎ、機種シートはすべて同じ枠組みで
生成し直す。行の挿入でセル結合がずれる崩れを避けるため、機種シートは手で直さない。

使い方:
    python3 build_param_sheet.py <元のxlsx> <出力xlsx>
"""
import math
import sys
import unicodedata

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import column_index_from_string as ci
from openpyxl.utils import get_column_letter as cl

from models import COMMON, MODELS, REVISIONS

FONT = "ＭＳ Ｐゴシック"
LAST_COL = "BJ"
COL_WIDTH = 2.296875
COL_PX = 19  # 1 列の概算幅（px）。折り返し判定を安全側にするため小さめ

FILL_LABEL = PatternFill("solid", fgColor=Color(theme=3, tint=0.7999816888943144))
FILL_GROUP = PatternFill("solid", fgColor=Color(theme=3, tint=0.5999938962981048))
FILL_TITLE = PatternFill("solid", fgColor="FF17195E")
FILL_TODO = PatternFill("solid", fgColor="FFFFFF99")
FILL_CONNECT = PatternFill("solid", fgColor="FFCCFFFF")
FILL_SHUTDOWN = PatternFill("solid", fgColor="FFFFCC99")
FILL_OPTION = PatternFill("solid", fgColor=Color(theme=0, tint=-0.1499984740745262))

THIN = Side(style="thin")
MEDIUM = Side(style="medium")
DOUBLE = Side(style="double")

TODO_MARK = "要記入"


def text_px(text, size):
    """MS Pゴシックでの表示幅を概算する（全角 1.0、半角 0.55 文字幅）。"""
    if text is None:
        return 0
    units = 0.0
    for ch in str(text):
        units += 1.0 if unicodedata.east_asian_width(ch) in "WFA" else 0.55
    return units * size * 1.45


class SheetWriter:
    """1 枚の機種シートを上から順に書く。"""

    def __init__(self, ws):
        self.ws = ws
        self.row = 1
        self.row_lines = {}
        self.row_min = {}
        for i in range(1, ci(LAST_COL) + 1):
            ws.column_dimensions[cl(i)].width = COL_WIDTH

    # ---- 低レベル -------------------------------------------------------
    def cell(self, row, c1, c2=None, value=None, *, rows=1, kind="value",
             size=10, bold=False, halign="left", valign="center", wrap=True,
             fill=None, color=None):
        c2 = c2 or c1
        ws = self.ws
        top = ws[f"{c1}{row}"]
        if isinstance(value, str) and value.startswith(TODO_MARK):
            fill = fill or FILL_TODO
        if kind == "label":
            fill = fill or FILL_LABEL
        top.value = value
        top.font = Font(name=FONT, size=size, bold=bold,
                        color=color if color else None)
        top.alignment = Alignment(horizontal=halign, vertical=valign,
                                  wrap_text=wrap)
        for r in range(row, row + rows):
            for c in range(ci(c1), ci(c2) + 1):
                x = ws.cell(r, c)
                x.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
                if fill:
                    x.fill = fill
        if (c1 != c2) or rows > 1:
            ws.merge_cells(f"{c1}{row}:{c2}{row + rows - 1}")
        if wrap and rows == 1 and value is not None and not str(value).startswith("="):
            width = (ci(c2) - ci(c1) + 1) * COL_PX - 6
            lines = 0
            for part in str(value).split("\n"):
                lines += max(1, math.ceil(text_px(part, size) / max(width, 1)))
            self.row_lines[row] = max(self.row_lines.get(row, 1), lines)
        return top

    def outline(self, r1, r2, c1="A", c2=LAST_COL, header_rows=0):
        """範囲の外枠を太線、見出し行の下を二重線にする。"""
        ws = self.ws
        for r in range(r1, r2 + 1):
            for c in range(ci(c1), ci(c2) + 1):
                x = ws.cell(r, c)
                b = x.border
                left, right, top, bottom = b.left, b.right, b.top, b.bottom
                if c == ci(c1):
                    left = MEDIUM
                if c == ci(c2):
                    right = MEDIUM
                if r == r1:
                    top = MEDIUM
                if r == r2:
                    bottom = MEDIUM
                if header_rows and r == r1 + header_rows - 1:
                    bottom = DOUBLE
                if header_rows and r == r1 + header_rows:
                    top = DOUBLE
                x.border = Border(left=left, right=right, top=top, bottom=bottom)

    def note_row(self, r, text):
        """表の最下行に注記を 1 行置く。注記が無ければ何もしない。"""
        if not text:
            return r
        self.cell(r, "A", LAST_COL, text, size=9)
        return r + 1

    def height_of(self, r):
        lines = self.row_lines.get(r, 1)
        h = 14 * lines + 4 if lines > 1 else 18
        return max(h, self.row_min.get(r, 0))

    def fit_note(self, text, r1, r2, c1, c2, size=9):
        """結合した備考欄に文章が収まるよう、行 r1～r2 の高さを広げる。"""
        if not text:
            return
        width = (ci(c2) - ci(c1) + 1) * COL_PX - 6
        lines = sum(max(1, math.ceil(text_px(p, size) / width)) for p in str(text).split("\n"))
        need = lines * (size + 5) + 10
        have = sum(self.height_of(r) for r in range(r1, r2 + 1))
        if need > have:
            extra = (need - have) / (r2 - r1 + 1)
            for r in range(r1, r2 + 1):
                self.row_min[r] = self.height_of(r) + extra

    def finish_heights(self):
        for r in set(self.row_lines) | set(self.row_min):
            self.ws.row_dimensions[r].height = round(self.height_of(r), 1)

    # ---- 中レベル -------------------------------------------------------
    def title(self, text):
        ws = self.ws
        for c in range(2, ci(LAST_COL) + 1):
            ws.cell(2, c).fill = FILL_TITLE
        x = ws["B2"]
        x.value = text
        x.font = Font(name=FONT, size=20, color="FFFFFFFF")
        x.alignment = Alignment(vertical="center")
        ws.row_dimensions[2].height = 30
        self.row = 4

    def chapter(self, text):
        x = self.ws[f"A{self.row}"]
        x.value = text
        x.font = Font(name=FONT, size=12, bold=True)
        self.row += 1

    def gap(self, n=1):
        self.row += n

    def block(self, header, rows, note=None, *, note_col="AY", main_end="AX",
              header_cells=None):
        """見出し行＋本文行＋右端の備考欄で 1 ブロックを書く。

        rows は 1 行ごとのセル定義のリスト。セル定義は
        (開始列, 終了列, 値, 種別[, 縦に結合する行数]) のタプル。種別は "label" / "value" / "skip"。
        備考欄は本文行をまとめて 1 つに結合する。note=False なら備考欄を作らない。
        """
        r0 = self.row
        if note is False:
            self.cell(r0, "A", LAST_COL, header, kind="label", bold=True, size=11)
        else:
            self.cell(r0, "A", main_end, header, kind="label", bold=True, size=11)
            self.cell(r0, note_col, LAST_COL, "備考", bold=True, size=11)
        r = r0 + 1
        for spec in header_cells or []:
            for c1, c2, v in spec:
                self.cell(r, c1, c2, v, kind="label", size=9, halign="center")
            r += 1
        body_start = r
        for spec in rows:
            for item in spec:
                c1, c2, v = item[:3]
                kind = item[3] if len(item) > 3 else "value"
                span = item[4] if len(item) > 4 else 1
                if kind == "skip":
                    continue
                self.cell(r, c1, c2, v, kind=kind, rows=span)
            r += 1
        last = r - 1
        if note is not False:
            n = self.cell(body_start, note_col, LAST_COL, note or None,
                          rows=max(1, last - body_start + 1), size=9,
                          valign="top")
            n.alignment = Alignment(horizontal="left", vertical="top",
                                    wrap_text=True)
            self.fit_note(note, body_start, max(body_start, last), note_col, LAST_COL)
        self.outline(r0, last, header_rows=1 + len(header_cells or []))
        self.row = last + 1
        return r0, last


# ---- 部品 -----------------------------------------------------------------
L = ("A", "J")
S1, S2, S3, S4 = ("K", "T"), ("U", "AD"), ("AE", "AN"), ("AO", "AX")
P1V = ("K", "AD")  # 2 項目行の 1 つ目の値
P2L = ("AE", "AN")  # 2 項目行の 2 つ目の見出し
P2V = ("AO", "AX")
FULL = ("K", "AX")


def lab(c, v, span=1):
    return (c[0], c[1], v, "label", span)


def skip(c):
    return (c[0], c[1], None, "skip")


def val(c, v):
    return (c[0], c[1], v, "value")


def row_full(label, value):
    return [lab(L, label), val(FULL, value)]


def row_pair(l1, v1, l2, v2):
    return [lab(L, l1), val(P1V, v1), lab(P2L, l2), val(P2V, v2)]


def row_quad(label, values, kinds=("value",) * 4):
    out = [lab(L, label)]
    for c, v, k in zip((S1, S2, S3, S4), values, kinds):
        out.append((c[0], c[1], v, k))
    return out


def port_names(m, member):
    ge = [f"GE{member}/0/{i}" for i in range(1, m["ge"] + 1)]
    xge = [f"10GE{member}/0/{i}" for i in range(1, m["xge"] + 1)]
    return ge, xge


# ---- シート本体 ------------------------------------------------------------
def build_sheet(wb, m):
    ws = wb.create_sheet(m["sheet"])
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    w = SheetWriter(ws)
    d = COMMON
    w.title(f"基本(LinkSphere {m['sheet']})設定")

    # 基本情報
    w.chapter("● 基本情報")
    w.block("お客様情報", [
        row_full("会社名", None),
        row_full("設置先住所", None),
        row_full("システム名", None),
        row_full("使用用途", None),
    ], note=False)
    w.gap()

    # 機器構成情報
    w.chapter("● 機器構成情報")
    hw_rows = [
        row_full("製品名", f"LinkSphere {m['sheet']}"),
        row_full("ボードタイプ", m["sheet"]),
        row_full("固定ポート構成", m["port_summary"]),
        row_full("ソフトウェアバージョン", m["sw_version"]),
        row_full("システムソフトウェアファイル", m["sw_file"]),
        row_pair("電源構成", m["power"], "冷却方式", m["cooling"]),
        row_pair("PoE対応", m["poe_std"], "PoE対象ポート", m["poe_ports"]),
        row_pair("PoE給電能力", m["poe_budget"], "スタックポート", m["stack_port_hw"]),
        [lab(L, "SFPトランシーバー（別途手配品）"), lab(("K", "P"), "製品型番"),
         val(("Q", "AF"), None), lab(("AG", "AJ"), "個数"), val(("AK", "AX"), None)],
        [lab(L, "SFP+／DAC／AOC（別途手配品）"), lab(("K", "P"), "製品型番"),
         val(("Q", "AF"), None), lab(("AG", "AJ"), "個数"), val(("AK", "AX"), None)],
    ]
    r_hw, _ = w.block("ハードウェア構成情報", hw_rows, note=m["hw_note"])
    product_cell = f"K{r_hw + 1}"
    w.gap()

    # スタック構成情報
    members = ["メンバー1", "メンバー2", "メンバー3", "メンバー4"]
    stack_rows = [
        row_quad("項目", members, kinds=("label",) * 4),
        row_quad("メンバーID", ["1", "2", "3", "4"]),
        row_quad("想定ロール", [None] * 4),
        row_quad("筐体ラベル名", [None] * 4),
        row_quad("シリアル番号", [None] * 4),
        row_quad("スタック優先度", [d["stack_priority"]] * 4),
    ]
    if m["dedicated_stack"]:
        stack_rows.append(row_quad("スタックポート", [m["stack_port_cfg"]] * 4))
    else:
        stack_rows += [
            row_quad("スタックポート1（stack-port n/1）", [m["stack_port_cfg"]] * 4),
            row_quad("スタックポート2（stack-port n/2）", [m["stack_port_cfg"]] * 4),
        ]
    stack_rows += [
        row_full("スタックドメイン", d["stack_domain"]),
        row_full("DAD（デュアルアクティブ検知）", d["stack_dad"]),
    ]
    w.block("スタック構成情報", stack_rows, note=d["stack_note"] + m.get("stack_note", ""))
    w.gap()

    # 基本設定
    w.chapter("● 基本設定")
    w.block("管理アクセス設定", [
        row_full("ホスト名（sysname）", None),
        row_pair("コンソール認証", d["console_auth"], "ユーザーインターフェース", d["console_ui"]),
        [lab(L, "ローカルユーザー", 3), lab(("K", "R"), "ユーザー名"), lab(("S", "Z"), "パスワード"),
         lab(("AA", "AH"), "権限レベル"), lab(("AI", "AX"), "サービスタイプ")],
        [skip(L), val(("K", "R"), None), val(("S", "Z"), None), val(("AA", "AH"), None), val(("AI", "AX"), None)],
        [skip(L), val(("K", "R"), None), val(("S", "Z"), None), val(("AA", "AH"), None), val(("AI", "AX"), None)],
        row_pair("VTY", d["vty"], "認証方式／許可プロトコル", d["vty_auth"]),
        row_pair("SSH（STelnet）サーバー", d["stelnet"], "SSH受付インターフェース", d["ssh_source"]),
        [lab(L, "SSH"), lab(("K", "R"), "RSA最小鍵長"), val(("S", "Z"), d["rsa_min"]),
         lab(("AA", "AD"), "Version"), val(("AE", "AH"), d["ssh_ver"]),
         lab(("AI", "AN"), "Timeout"), val(("AO", "AQ"), d["ssh_timeout"]),
         lab(("AR", "AU"), "認証リトライ"), val(("AV", "AX"), d["ssh_retry"])],
        [lab(L, "管理アクセスACL"), val(S1, d["acl_ssh"]), val(S2, d["acl_web"]),
         val(S3, d["acl_snmp"]), val(S4, None)],
    ], note=d["mgmt_note"])
    w.gap()

    w.block("時刻設定", [
        row_full("タイムゾーン", d["timezone"]),
        row_pair("NTPサーバー1", d["ntp"], "NTPサーバー2", d["ntp"]),
    ], note=d["time_note"])
    w.gap()

    w.block("DNS設定", [
        row_pair("DNS名前解決（dns resolve）", d["dns_resolve"], "送信元IP（dns server source-ip）", None),
        row_pair("DNSサーバー1", None, "DNSサーバー2", None),
    ], note=d["dns_note"])
    w.gap()

    w.block("Syslog設定", [
        row_pair("Syslogサーバー1", d["syslog"], "Syslogサーバー2", d["syslog"]),
        row_pair("Severity Level", d["syslog_level"], "Facility", d["syslog_facility"]),
    ], note=d["syslog_note"])
    w.gap()

    w.block("MTU設定", [
        row_pair("最大フレーム長（jumboframe enable）", d["jumbo"], "Vlanif の MTU（mtu）", d["vlanif_mtu"]),
        [lab(L, "MTU を変える Vlanif", 3), lab(("K", "AD"), "Vlanif"), lab(("AE", "AX"), "MTU（バイト）")],
        [skip(L), val(("K", "AD"), None), val(("AE", "AX"), None)],
        [skip(L), val(("K", "AD"), None), val(("AE", "AX"), None)],
    ], note=d["mtu_note"])
    w.gap()

    w.block("SNMP設定（初期状態：Agent無効）", [
        row_pair("SNMP Agent", d["snmp_agent"], "Version", d["snmp_version"]),
        [lab(L, "SNMPv3ユーザー", 2), lab(("K", "R"), "ユーザー名"), lab(("S", "Z"), "認証方式"),
         lab(("AA", "AH"), "暗号方式"), lab(("AI", "AP"), "セキュリティレベル"), lab(("AQ", "AX"), "グループ")],
        [skip(L), val(("K", "R"), None), val(("S", "Z"), d["snmp_auth"]), val(("AA", "AH"), d["snmp_priv"]),
         val(("AI", "AP"), d["snmp_level"]), val(("AQ", "AX"), None)],
        row_full("SNMPアクセス制限（ACL）", d["acl_snmp"]),
        [lab(L, "Trapサーバー"), lab(("K", "T"), "IP Address"), lab(("U", "X"), "Version"),
         lab(("Y", "AF"), "Trap送信元IF"), lab(("AG", "AX"), "Security Name（v3ユーザー名）")],
        [lab(L, "　Trapサーバー1"), val(("K", "T"), None), val(("U", "X"), None), val(("Y", "AF"), None), val(("AG", "AX"), None)],
        [lab(L, "　Trapサーバー2"), val(("K", "T"), None), val(("U", "X"), None), val(("Y", "AF"), None), val(("AG", "AX"), None)],
    ], note=d["snmp_note"])
    w.gap()

    radius_server_cols = (("K", "V"), ("W", "Z"), ("AA", "AX"))
    w.block("RADIUS設定（初期状態：未設定）", [
        row_full("サーバーテンプレート名（radius-server template）", None),
        [lab(L, "RADIUSサーバー"), lab(radius_server_cols[0], "IP Address"), lab(radius_server_cols[1], "Port"),
         lab(radius_server_cols[2], "共有鍵")],
    ] + [[lab(L, f"　{name}")] + [val(c, None) for c in radius_server_cols]
         for name in ("認証サーバー1", "認証サーバー2", "アカウンティングサーバー1", "アカウンティングサーバー2")] + [
        row_pair("認証の順序（authentication-mode）", None, "アカウンティング（accounting-mode）", None),
        row_full("適用するユーザー（ドメイン）", None),
    ], note=d["radius_note"])
    w.gap()

    w.block("STP設定", [
        row_pair("STPモード", d["stp_mode"], "ブリッジ優先度", d["stp_priority"]),
        row_pair("Hello／Max Age", d["stp_hello"], "Forward Delay／Max Hops", d["stp_fwd"]),
        row_full("BPDU保護（グローバル）", d["bpdu_protection"]),
        row_pair("MSTリージョン名", d["region_name"], "Revision", d["region_rev"]),
        row_full("VLANマッピング", d["region_map"]),
        row_pair("ポート優先度", d["port_priority"], "パスコスト", d["path_cost"]),
        row_full("エッジポート", d["edged_default"]),
    ], note=d["stp_note"])
    w.gap()

    w.block("LLDP設定", [
        row_full("LLDP", d["lldp"]),
    ], note=d["lldp_note"])
    w.gap()

    w.block("管理用インターフェース", [
        [lab(L, "MEth0/0/0（論理IF）", 2), lab(S1, "モード"), lab(S2, "MTU"), lab(S3, "IP Address"), lab(S4, "Subnet Mask")],
        [skip(L), val(S1, "Route Port"), val(S2, "1500"), val(S3, "未設定"), val(S4, "未設定")],
        [lab(L, "管理用Vlanif", 2), lab(S1, "VLAN ID"), lab(S2, "IP Address"), lab(S3, "Subnet Mask"), lab(S4, "Default Gateway")],
        [skip(L), val(S1, None), val(S2, None), val(S3, None), val(S4, None)],
    ], note=d["mgmt_if_note"])
    w.gap()

    w.block("システム初期設定", [[val(("A", "AX"), line)] for line in d["system_defaults"]],
            note="出荷時の既定で有効・無効になっている設定。")
    w.gap()

    w.block("バナー設定", [
        row_full("ログイン前（header login information）", None),
        row_full("ログイン後（header shell information）", None),
    ], note=None)
    w.gap()

    w.block("VLAN設定", [
        [lab(("A", "C"), "ID"), lab(("D", "W"), "Name"), lab(("X", "Z"), "ID"), lab(("AA", "AX"), "Name")],
    ] + [[val(("A", "C"), None), val(("D", "W"), None), val(("X", "Z"), None), val(("AA", "AX"), None)]
         for _ in range(6)], note=d["vlan_note"])
    w.gap()

    l3_header = [
        [("A", "F", "ID／IF"), ("G", "L", "IP Address"), ("M", "R", "Subnet Mask"), ("S", "Z", "Description"),
         ("AA", "AH", "ACL（traffic-filter）"), ("AI", LAST_COL, "VRRP")],
        [("A", "F", ""), ("G", "L", ""), ("M", "R", ""), ("S", "Z", ""),
         ("AA", "AE", "番号／名前"), ("AF", "AH", "方向"), ("AI", "AK", "VRID"), ("AL", "AQ", "仮想IP"),
         ("AR", "AT", "優先度"), ("AU", "AX", "通知間隔"), ("AY", "BC", "プリエンプト遅延"),
         ("BD", LAST_COL, "監視IF／優先度減算")],
    ]
    l3_cols = [("A", "F"), ("G", "L"), ("M", "R"), ("S", "Z"), ("AA", "AE"), ("AF", "AH"),
               ("AI", "AK"), ("AL", "AQ"), ("AR", "AT"), ("AU", "AX"), ("AY", "BC"), ("BD", LAST_COL)]
    for title_text in ("Vlanif設定", "Routed Port設定"):
        w.l3_table(title_text, l3_header, l3_cols, 5)
    w.block("備考", [[val(("A", LAST_COL), d["vrrp_note"])]], note=False)
    w.gap()

    w.block("DHCP設定（初期状態：未設定）", [
        row_full("DHCP機能（dhcp enable）", d["dhcp_enable"]),
        [lab(L, "DHCPリレー", 4), lab(("K", "T"), "Vlanif"), lab(("U", "AJ"), "DHCPサーバー1"), lab(("AK", "AX"), "DHCPサーバー2")],
    ] + [[skip(L), val(("K", "T"), None), val(("U", "AJ"), None), val(("AK", "AX"), None)] for _ in range(3)] + [
        [lab(L, "DHCPサーバー（アドレスプール）", 3), lab(("K", "R"), "プール名"), lab(("S", "AB"), "ネットワーク／マスク"),
         lab(("AC", "AL"), "ゲートウェイ"), lab(("AM", "AX"), "DNSサーバー")],
    ] + [[skip(L), val(("K", "R"), None), val(("S", "AB"), None), val(("AC", "AL"), None), val(("AM", "AX"), None)]
         for _ in range(2)] + [
        row_pair("DHCPスヌーピング", d["dhcp_snooping"], "信頼ポート（trusted）", None),
    ], note=d["dhcp_note"])
    w.gap()

    # ポート構成情報
    w.port_diagram(m, product_cell)
    w.gap()

    # ポート設定
    w.port_table(m)
    w.gap()

    if m["poe_std"] not in ("非対応", None):
        w.block("PoE設定", [
            row_pair("給電（全対象ポート）", d["poe_enable"], "給電優先度", d["poe_priority"]),
            row_full("ポート上限電力", d["poe_power"]),
            [lab(L, "個別設定", 5), lab(("K", "R"), "Interface"), lab(("S", "Z"), "給電"),
             lab(("AA", "AH"), "給電優先度"), lab(("AI", "AX"), "上限電力（mW）")],
        ] + [[skip(L), val(("K", "R"), None), val(("S", "Z"), None), val(("AA", "AH"), None), val(("AI", "AX"), None)]
             for _ in range(4)], note=d["poe_note"])
        w.gap()

    w.trunk_table()
    w.gap()

    # ルーティング
    w.block("ルーティング設定", [
        [lab(L, "静的ルート", 5), lab(S1, "宛先ネットワーク"), lab(S2, "Subnet Mask"), lab(S3, "Next Hop"),
         lab(S4, f"Preference（既定 {d['static_pref']}）")],
    ] + [[skip(L), val(S1, None), val(S2, None), val(S3, None), val(S4, None)] for _ in range(4)] + [
        row_pair("動的ルーティング", None, "Process ID", None),
        row_pair("Router-ID", None, "Area", None),
        row_pair("network（ネットワーク）", None, "ワイルドカードマスク", None),
        row_pair("集約（Summary）", None, "Graceful Restart", None),
        row_pair("silent-interface", None, "import-route", None),
    ], note=d["routing_note"])
    w.gap()

    w.acl_table()
    w.gap()

    w.qos_tables()
    w.gap()

    w.storm_table()

    w.finish_heights()
    last = w.row
    ws.print_area = f"A1:{LAST_COL}{last}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "&P / &N"
    return ws


# ---- 表の書き出し（SheetWriter へ追加） ----------------------------------
def l3_table(self, title_text, header, cols, nrows):
    r0 = self.row
    self.cell(r0, "A", LAST_COL, title_text, kind="label", bold=True, size=11)
    r = r0 + 1
    # 2 段見出し: 1 段目で縦結合するものは 2 段分にする
    top, bottom = header
    for (c1, c2, v), (b1, b2, bv) in zip(top[:4], bottom[:4]):
        self.cell(r, c1, c2, v, kind="label", rows=2, size=9, halign="center")
    self.cell(r, "AA", "AH", top[4][2], kind="label", size=9, halign="center")
    self.cell(r, "AI", LAST_COL, top[5][2], kind="label", size=9, halign="center")
    for c1, c2, v in bottom[4:]:
        self.cell(r + 1, c1, c2, v, kind="label", size=9, halign="center")
    r += 2
    for _ in range(nrows):
        for c1, c2 in cols:
            self.cell(r, c1, c2, None)
        r += 1
    self.outline(r0, r - 1, header_rows=3)
    self.row = r


def port_diagram(self, m, product_cell):
    ws = self.ws
    r0 = self.row
    self.cell(r0, "A", LAST_COL, "ポート構成情報（メンバー1。メンバー2～4も同じ配置）",
              kind="label", bold=True, size=11)
    r_top = r0 + 2
    r_bot = r0 + 5
    step = 4 if m["ge"] <= 24 else 2

    def box(r, col, name, num):
        c1 = cl(col)
        c2 = cl(col + 1)
        self.cell(r, c1, c2, name, size=8, halign="center", wrap=False)
        self.cell(r + 1, c1, c2, num, size=7, halign="center", wrap=False)
        for c in (col, col + 1):
            ws.cell(r, c).border = Border(left=THIN if c == col else None,
                                          right=THIN if c == col + 1 else None, top=THIN)
            ws.cell(r + 1, c).border = Border(left=THIN if c == col else None,
                                              right=THIN if c == col + 1 else None, bottom=THIN)

    for i in range(1, m["ge"] + 1):
        pair = (i + 1) // 2
        col = 2 + (pair - 1) * step
        r = r_top if i % 2 == 1 else r_bot
        box(r, col, "GE", f"1/0/{i}")
    for i in range(1, m["xge"] + 1):
        col = ci("BA") + (i - 1) * 2
        box(r_top, col, "10GE", f"1/0/{i}")
    if m["dedicated_stack"]:
        for i in range(1, 3):
            col = ci("BA") + (i - 1) * 2
            box(r_bot, col, "STACK", str(i))
    r_leg = r_bot + 3
    self.cell(r_leg, "C", "J", "Connect", size=10, fill=FILL_CONNECT)
    self.cell(r_leg, "M", "T", "Not Connect", size=10)
    self.cell(r_leg, "W", "AD", "Shutdown", size=10, fill=FILL_SHUTDOWN)
    self.cell(r_leg, "AG", "AN", "Option", size=10, fill=FILL_OPTION)
    self.cell(r_leg, "AP", "BG", f"={product_cell}", size=11, halign="center")
    for c in range(ci("AP"), ci("BG") + 1):
        ws.cell(r_leg, c).border = Border(bottom=MEDIUM)
    r_note = r_leg + 2
    self.cell(r_note, "A", LAST_COL, "備考", kind="label", bold=True, size=11)
    note = m["port_summary"] + "。" + m.get("diagram_note", "")
    self.cell(r_note + 1, "A", LAST_COL, note, size=10)
    self.outline(r0, r_note - 1, header_rows=1)
    self.outline(r_note, r_note + 1, header_rows=1)
    self.row = r_note + 2


PORT_COLS = [
    ("A", "F", "Interface"), ("G", "I", "Media"), ("J", "L", "Speed"), ("M", "O", "Duplex"),
    ("P", "R", "Port\nStatus"), ("S", "X", "link-type"), ("Y", "AA", "PVID"), ("AB", "AF", "許可VLAN"),
    ("AG", "AI", "Eth-Trunk"), ("AJ", "AM", "エッジ\nポート"), ("AN", "AQ", "BPDU\nFilter"),
    ("AR", "AU", "ループ\n保護"), ("AV", "AY", "Up/Down\nTrap"), ("AZ", LAST_COL, "Description"),
]
PORT_SUB = ("S", "Y", "AB", "AJ", "AN", "AR")  # グループ見出しの下に並ぶ列


def port_table(self, m):
    d = COMMON
    r0 = self.row
    self.cell(r0, "A", LAST_COL, "ポート設定", kind="label", bold=True, size=11)
    r = r0 + 1
    groups = {"S": ("S", "AF", "VLAN"), "AJ": ("AJ", "AU", "STP")}
    for c1, c2, v in PORT_COLS:
        if c1 in PORT_SUB:
            continue
        self.cell(r, c1, c2, v, kind="label", rows=2, size=9, halign="center")
    for c1, c2, v in groups.values():
        self.cell(r, c1, c2, v, kind="label", size=9, halign="center")
    for c1, c2, v in PORT_COLS:
        if c1 in PORT_SUB:
            self.cell(r + 1, c1, c2, v, kind="label", size=9, halign="center")
    self.row_lines[r + 1] = 2
    r += 2
    header_rows = 3
    for member in range(1, 5):
        label = f"メンバー{member}" + ("" if member == 1 else "（スタック構成時のみ記入）")
        self.cell(r, "A", LAST_COL, label, fill=FILL_GROUP, bold=True, size=10)
        r += 1
        ge, xge = port_names(m, member)
        for name in ge + xge:
            is_x = name.startswith("10GE")
            vals = {
                "A": name, "G": "SFP+" if is_x else "RJ45",
                "J": d["speed_x"] if is_x else d["speed_ge"],
                "M": d["duplex_x"] if is_x else d["duplex_ge"],
                "P": "有効", "S": d["link_type"], "Y": "1", "AB": d["allow_vlan"],
                "AG": "-", "AJ": "disabled", "AN": "disabled", "AR": "disabled",
                "AV": d["updown_trap"], "AZ": None,
            }
            for c1, c2, _ in PORT_COLS:
                self.cell(r, c1, c2, vals[c1], size=8 if c1 == "S" else 9, wrap=False)
            r += 1
    r = self.note_row(r, d["port_note"])
    self.outline(r0, r - 1, header_rows=header_rows)
    self.row = r


TRUNK_COLS = [
    ("A", "F", "Eth-Trunk"), ("G", "K", "動作モード\n（mode）"), ("L", "T", "メンバーポート"), ("U", "X", "link-type"),
    ("Y", "AA", "PVID"), ("AB", "AG", "許可VLAN"), ("AH", "AJ", "Port\nStatus"), ("AK", "AN", "エッジ\nポート"),
    ("AO", "AR", "BPDU\nFilter"), ("AS", "AV", "ループ\n保護"), ("AW", "AZ", "Up/Down\nTrap"), ("BA", LAST_COL, "Description"),
]


def trunk_table(self):
    d = COMMON
    r0 = self.row
    self.cell(r0, "A", LAST_COL, "Eth-Trunk設定", kind="label", bold=True, size=11)
    self.cell(r0 + 1, "A", "J", "負荷分散方式（load-balance）", kind="label", size=10)
    self.cell(r0 + 1, "K", LAST_COL, d["trunk_lb"], size=10)
    r = r0 + 2
    for c1, c2, v in TRUNK_COLS:
        self.cell(r, c1, c2, v, kind="label", size=9, halign="center")
    self.row_lines[r] = 2
    r += 1
    for i in range(5):
        for c1, c2, _ in TRUNK_COLS:
            v = {"A": "Eth-Trunk", "G": None, "AW": d["updown_trap"]}.get(c1)
            self.cell(r, c1, c2, v, size=9, wrap=False)
        r += 1
    r = self.note_row(r, d["trunk_note"])
    self.outline(r0, r - 1, header_rows=3)
    self.row = r


ACL_COLS = [
    ("A", "C", "No"), ("D", "J", "ACL番号／名前"), ("K", "M", "Rule ID"), ("N", "P", "Action"),
    ("Q", "S", "Protocol"), ("T", "Y", "送信元IP"), ("Z", "AE", "ワイルドカード"), ("AF", "AH", "Port"),
    ("AI", "AN", "宛先IP"), ("AO", "AT", "ワイルドカード"), ("AU", "AW", "Port"), ("AX", "AZ", "logging"),
    ("BA", LAST_COL, "備考"),
]


def acl_table(self):
    d = COMMON
    r0 = self.row
    self.cell(r0, "A", LAST_COL, "ACL設定", kind="label", bold=True, size=11)
    r = r0 + 1
    for c1, c2, v in ACL_COLS:
        self.cell(r, c1, c2, v, kind="label", size=9, halign="center")
    r += 1
    for _ in range(6):
        for c1, c2, _v in ACL_COLS:
            self.cell(r, c1, c2, None, size=9)
        r += 1
    r = self.note_row(r, d["acl_note"])
    self.outline(r0, r - 1, header_rows=2)
    self.row = r


def qos_tables(self):
    d = COMMON
    r0 = self.row
    self.cell(r0, "A", LAST_COL, "QoS設定（オプション／初期状態：未設定）", kind="label", bold=True, size=11)
    r = r0 + 1
    tables = [
        ("流分類（traffic classifier）", [("K", "T", "Classifier名"), ("U", "Z", "and／or"), ("AA", LAST_COL, "マッチ条件（if-match）")]),
        ("流動作（traffic behavior）", [("K", "T", "Behavior名"), ("U", LAST_COL, "動作（remark／car／deny 等）")]),
        ("流ポリシー（traffic policy）", [("K", "T", "Policy名"), ("U", "AJ", "Classifier名"), ("AK", LAST_COL, "Behavior名")]),
        ("適用インターフェース", [("K", "T", "Interface"), ("U", "Z", "方向"), ("AA", "AH", "trust"),
                            ("AI", "AR", "スケジュール（PQ／DRR）"), ("AS", LAST_COL, "DRR重み")]),
    ]
    for title_text, cols in tables:
        self.cell(r, "A", "J", title_text, kind="label", rows=4, size=10, valign="top")
        for c1, c2, v in cols:
            self.cell(r, c1, c2, v, kind="label", size=9, halign="center")
        r += 1
        for _ in range(3):
            for c1, c2, _v in cols:
                self.cell(r, c1, c2, None, size=9)
            r += 1
    r = self.note_row(r, d["qos_note"])
    self.outline(r0, r - 1, header_rows=1)
    self.row = r


STORM_COLS = [
    ("A", "J", "Interface"), ("K", "P", "対象パケット"), ("Q", "U", "単位"), ("V", "Y", "min-rate"),
    ("Z", "AC", "max-rate"), ("AD", "AH", "action"), ("AI", "AK", "log"), ("AL", "AN", "trap"),
    ("AO", "AR", "interval"), ("AS", LAST_COL, "備考"),
]


def storm_table(self):
    d = COMMON
    r0 = self.row
    self.cell(r0, "A", LAST_COL, "ストーム制御設定（storm control。オプション／初期状態：未設定）", kind="label", bold=True, size=11)
    r = r0 + 1
    for c1, c2, v in STORM_COLS:
        self.cell(r, c1, c2, v, kind="label", size=9, halign="center")
    r += 1
    for _ in range(4):
        for c1, c2, _v in STORM_COLS:
            self.cell(r, c1, c2, None, size=9)
        r += 1
    self.cell(r, "A", "J", "自動復旧（error-down 時）", kind="label", size=10)
    self.cell(r, "K", LAST_COL, d["storm_recovery"], size=10)
    r += 1
    r = self.note_row(r, d["storm_note"])
    self.outline(r0, r - 1, header_rows=2)
    self.row = r


SheetWriter.l3_table = l3_table
SheetWriter.port_diagram = port_diagram
SheetWriter.port_table = port_table
SheetWriter.trunk_table = trunk_table
SheetWriter.acl_table = acl_table
SheetWriter.qos_tables = qos_tables
SheetWriter.storm_table = storm_table


def update_revision(ws):
    """改訂履歴の空き行へ、REVISIONS の版を順に足す。既存の行は書き換えない。"""
    for r in range(6, 19):
        ws.row_dimensions[r].height = None
    for rev in REVISIONS:
        row = None
        for r in range(6, 19):
            if all(ws.cell(r, c).value in (None, "") for c in (1, 6, 14, 18)):
                row = r
                break
        if row is None:
            raise RuntimeError("改訂履歴に空き行がありません")
        ws.cell(row, 1).value = rev["version"]
        ws.cell(row, 6).value = rev["date"]
        ws.cell(row, 14).value = rev["kind"]
        ws.cell(row, 18).value = rev["summary"]
        ws.cell(row, 18).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 14 * rev["summary"].count("\n") + 30
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main(src, dst):
    wb = load_workbook(src)
    for name in [s.title for s in wb.worksheets if s.title != "改訂履歴"]:
        wb.remove(wb[name])
    update_revision(wb["改訂履歴"])
    for m in MODELS:
        build_sheet(wb, m)
    wb.active = 0
    wb.save(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2])
